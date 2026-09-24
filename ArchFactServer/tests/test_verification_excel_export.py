from io import BytesIO

from openpyxl import load_workbook

from app.services.verification_excel_export import build_machine_verification_excel


def test_full_machine_verification_excel_contains_summary_and_detail_rows() -> None:
    content = build_machine_verification_excel(
        version={
            "version": 1,
            "experiment_name": "实验 E1",
            "matching_version_id": "M0",
            "assertion_baseline": {"name": "LLM 断言 V1"},
            "rules": [{"id": 1, "title": "ID", "description": "唯一", "enabled": True}],
            "report": {
                "total_artifacts": 1,
                "full_pass_count": 0,
                "full_fail_count": 1,
                "full_uncertain_count": 0,
                "sample_count": 1,
                "reviewed_count": 1,
                "field_error_distribution": [{"key": "artifact_id_missing", "count": 1}],
            },
            "items": [
                {
                    "record_id": "record-1",
                    "verdict": "failed",
                    "machine_verdict": "failed",
                    "consensus_status": "agreed",
                }
            ],
        },
        machine_run={"_id": "run-1"},
        machine_items=[
            {
                "record_id": "record-1",
                "verdict": "failed",
                "confidence": 0.8,
                "reason": "编号缺失",
                "field_results": [
                    {
                        "field": "artifact_id",
                        "verdict": "failed",
                        "failure_code": "artifact_id_missing",
                        "causes_overall_failure": True,
                        "reason": "未提取到编号",
                    }
                ],
            }
        ],
        records=[
            {
                "_id": "record-1",
                "source_pages": [3],
                "thumbnail_region_id": "crop-1",
                "fields": {"artifact_id": {"value": "M1:1"}},
            }
        ],
    )

    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["校验汇总", "全量校验明细", "字段判断明细", "人工审核样本", "断言规则"]
    assert workbook["全量校验明细"].max_row == 2
    assert workbook["全量校验明细"]["C2"].value == "M1:1"
    assert workbook["字段判断明细"].max_row == 2
