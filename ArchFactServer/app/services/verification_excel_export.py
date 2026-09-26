"""Excel export for a persisted full machine-verification version."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="B96420")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="803F14")
_SUBHEADER_FILL = PatternFill("solid", fgColor="F8E8D9")
_FAILURE_LABELS = {
    "artifact_id_missing": "器物编号缺失",
    "artifact_id_duplicate": "器物编号重复",
    "artifact_id_evidence_conflict": "器物编号与证据冲突",
    "sequence_crop_conflict": "序号与裁剪图匹配错误",
    "figure_caption_missing": "图注缺失",
    "figure_caption_evidence_conflict": "图注与证据冲突",
    "color_plate_relation_conflict": "彩图关联冲突",
    "artifact_crop_missing": "器物裁剪图缺失",
    "structured_measurements_error": "尺寸结构化字段错误",
    "structured_classification_error": "分类结构化字段错误",
    "structured_field_evidence_conflict": "结构化字段与证据冲突",
    "text_evidence_conflict": "文本证据冲突",
    "unclassified_failure": "未分类错误",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    value = str(value).replace("\x00", "")[:32_000]
    # Keep OCR/LLM content as literal text rather than an Excel formula.
    return f"'{value}" if value[:1] in {"=", "+", "-", "@"} else value


def _field_value(record: dict[str, Any], key: str) -> str:
    field = (record.get("fields") or {}).get(key) or {}
    return _text(field.get("value", field.get("raw_value")))


def _artifact_id(record: dict[str, Any]) -> str:
    return _field_value(record, "artifact_id") or _text(record.get("artifact_id"))


def _field_failure_summary(field_results: list[dict[str, Any]]) -> tuple[str, str]:
    failed = [
        result
        for result in field_results
        if result.get("verdict") == "failed" and result.get("causes_overall_failure", True)
    ]
    fields = "、".join(_text(result.get("field")) for result in failed if result.get("field"))
    codes = "、".join(_text(result.get("failure_code")) for result in failed if result.get("failure_code"))
    return fields, codes


def _style_sheet(sheet, *, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        longest = max(
            (len(_text(sheet.cell(row=row, column=column).value)) for row in range(1, min(sheet.max_row, 150) + 1)),
            default=8,
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, 11), 42)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_machine_verification_excel(
    *,
    version: dict[str, Any],
    machine_run: dict[str, Any] | None,
    machine_items: list[dict[str, Any]],
    records: list[dict[str, Any]],
) -> bytes:
    """Create a self-contained workbook without writing temporary files."""

    workbook = Workbook()
    summary = workbook.active
    summary.title = "校验汇总"
    report = version.get("report") or {}
    baseline = version.get("assertion_baseline") or {}
    summary.append(["全量机器校验导出"])
    summary["A1"].font = _TITLE_FONT
    summary.merge_cells("A1:D1")
    summary.append(["项目", _text(version.get("experiment_name")), "断言版本", f"V{version.get('version', '')}"])
    summary.append(["断言基准", _text(baseline.get("name")), "匹配版本", _text(version.get("matching_version_id"))])
    summary.append(["全量器物", report.get("total_artifacts", 0), "机器通过", report.get("full_pass_count", 0)])
    summary.append(["机器错误", report.get("full_fail_count", 0), "机器不确定", report.get("full_uncertain_count", 0)])
    summary.append(["样本数量", report.get("sample_count", 0), "人工已审核", report.get("reviewed_count", 0)])
    summary.append(["错误覆盖率", report.get("error_coverage"), "断言精准率", report.get("error_precision")])
    summary.append(["人机一致性", report.get("human_machine_alignment"), "人工复核负担", report.get("review_load")])
    summary.append(["模型不可用条数", report.get("model_unavailable_count", 0), "机器任务 ID", _text((machine_run or {}).get("_id"))])
    summary.append([])
    summary.append(["错误字段", "错误代码", "数量"])
    for item in report.get("field_error_distribution") or []:
        key = _text(item.get("key"))
        summary.append([_FAILURE_LABELS.get(key, key), key, item.get("count", 0)])
    for cell in summary[11]:
        cell.fill = _SUBHEADER_FILL
        cell.font = Font(bold=True, color="803F14")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 28
    summary.column_dimensions["C"].width = 20
    summary.column_dimensions["D"].width = 32
    summary.sheet_view.showGridLines = False
    for row in range(2, 10):
        for column in (2, 4):
            summary.cell(row=row, column=column).alignment = Alignment(vertical="top", wrap_text=True)

    by_record_id = {str(record.get("_id")): record for record in records}
    details = workbook.create_sheet("全量校验明细")
    details.append([
        "序号", "器物记录 ID", "器物编号", "来源页", "类别", "质地", "表面颜色", "尺寸",
        "形态描述", "图注", "裁剪图区域 ID", "机器结论", "置信度", "模型服务异常",
        "判定原因", "触发字段", "错误代码", "字段判断 JSON",
    ])
    for index, item in enumerate(machine_items, start=1):
        record = by_record_id.get(str(item.get("record_id")), {})
        field_results = item.get("field_results") or []
        triggered_fields, failure_codes = _field_failure_summary(field_results)
        details.append([
            index,
            _text(item.get("record_id")),
            _artifact_id(record),
            "、".join(str(page) for page in record.get("source_pages") or []),
            _field_value(record, "category"),
            _field_value(record, "texture") or _field_value(record, "material"),
            _field_value(record, "surface_color"),
            _field_value(record, "measurements"),
            _field_value(record, "morphological_description"),
            _field_value(record, "figure_caption"),
            _text(record.get("primary_artifact_region_id") or record.get("thumbnail_region_id")),
            _text(item.get("verdict")),
            item.get("confidence"),
            "是" if item.get("model_unavailable") else "否",
            _text(item.get("reason")),
            triggered_fields,
            failure_codes,
            _text(field_results),
        ])
    _style_sheet(details)

    fields = workbook.create_sheet("字段判断明细")
    fields.append([
        "器物记录 ID", "器物编号", "机器结论", "字段", "字段结论", "错误代码",
        "导致整体失败", "字段原因", "字段判断 JSON",
    ])
    for item in machine_items:
        record = by_record_id.get(str(item.get("record_id")), {})
        for result in item.get("field_results") or []:
            fields.append([
                _text(item.get("record_id")),
                _artifact_id(record),
                _text(item.get("verdict")),
                _text(result.get("field")),
                _text(result.get("verdict")),
                _text(result.get("failure_code")),
                "是" if result.get("causes_overall_failure") else "否",
                _text(result.get("reason")),
                _text(result),
            ])
    _style_sheet(fields)

    samples = workbook.create_sheet("人工审核样本")
    samples.append([
        "样本序号", "器物记录 ID", "器物编号", "人工结论", "人工错误代码", "人工说明",
        "机器结论", "机器置信度", "机器原因", "AI 复核结论", "AI 复核置信度", "AI 复核原因",
        "人机一致状态",
    ])
    for index, item in enumerate(version.get("items") or [], start=1):
        record = by_record_id.get(str(item.get("record_id")), {})
        samples.append([
            index,
            _text(item.get("record_id")),
            _artifact_id(record),
            _text(item.get("verdict")),
            _text(item.get("failure_code")),
            _text(item.get("failure_reason")),
            _text(item.get("machine_verdict")),
            item.get("machine_confidence"),
            _text(item.get("machine_reason")),
            _text(item.get("ai_verdict")),
            item.get("ai_confidence"),
            _text(item.get("ai_reason")),
            _text(item.get("consensus_status")),
        ])
    _style_sheet(samples)

    rules = workbook.create_sheet("断言规则")
    rules.append(["来源", "规则 ID", "启用", "规则名称", "规则说明"])
    for rule in version.get("rules") or []:
        rules.append([
            _text(rule.get("source") or "selected"),
            _text(rule.get("id")),
            "是" if rule.get("enabled", True) else "否",
            _text(rule.get("title")),
            _text(rule.get("description")),
        ])
    _style_sheet(rules)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float) and 0 <= cell.value <= 1:
                    cell.number_format = "0.00%"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
